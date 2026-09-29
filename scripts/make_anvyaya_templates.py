"""
Generates the blank Excel templates for the Scripture Anvyayas section:

  content/anvyaya/sb/sb-01.xlsx … sb-09.xlsx, sb-10-part-1 … part-4, sb-11, sb-12
  content/anvyaya/bg/bg-01.xlsx … bg-18.xlsx

Only needed if you want a fresh blank file. Existing files are NOT
overwritten unless you pass --force (so your typed-in data is safe).

Usage:  python3 scripts/make_anvyaya_templates.py [--force]
Needs:  pip install openpyxl
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "S.N",
    "Chapter No",
    "Chapter Name",
    "File Number",
    "Verse No",
    "IAST-Verse",
    "Devanagari-Verse",
    "Vedabase Link",
    "Anvyaya",
    "English Translation",
]
WIDTHS = [6, 11, 30, 12, 11, 46, 46, 42, 50, 50]
TEXT_COLS = [2, 5]  # "Chapter No" and "Verse No" are kept as TEXT so 18.10 never becomes 18.1

ROOT = os.path.join(os.path.dirname(__file__), "..", "content", "anvyaya")
FORCE = "--force" in sys.argv

SB_FILES = (
    [(f"sb-{n:02d}", n) for n in range(1, 10)]
    + [(f"sb-10-part-{p}", 10) for p in range(1, 5)]
    + [("sb-11", 11), ("sb-12", 12)]
)
BG_FILES = [(f"bg-{n:02d}", n) for n in range(1, 19)]

SAMPLES = {
    "sb-01": [
        1, 1, "Questions by the Sages", 1, "1.1.1",
        "janmādy asya yato 'nvayād itarataś cārtheṣv abhijñaḥ svarāṭ / "
        "tene brahma hṛdā ya ādi-kavaye muhyanti yat sūrayaḥ / "
        "tejo-vāri-mṛdāṁ yathā vinimayo yatra tri-sargo 'mṛṣā / "
        "dhāmnā svena sadā nirasta-kuhakaṁ satyaṁ paraṁ dhīmahi",
        "जन्माद्यस्य यतोऽन्वयादितरतश्चार्थेष्वभिज्ञः स्वराट् / "
        "तेने ब्रह्म हृदा य आदिकवये मुह्यन्ति यत्सूरयः / "
        "तेजोवारिमृदां यथा विनिमयो यत्र त्रिसर्गोऽमृषा / "
        "धाम्ना स्वेन सदा निरस्तकुहकं सत्यं परं धीमहि",
        "https://vedabase.io/en/library/sb/1/1/1/",
        "SAMPLE ROW — delete it and type your own anvyaya here.",
        "SAMPLE ROW — delete it and type your own translation here.",
    ],
    "bg-01": [
        1, 1, "Observing the Armies on the Battlefield of Kurukṣetra", 1, "1.1",
        "dhṛtarāṣṭra uvāca / dharma-kṣetre kuru-kṣetre / samavetā yuyutsavaḥ / "
        "māmakāḥ pāṇḍavāś caiva / kim akurvata sañjaya",
        "धृतराष्ट्र उवाच / धर्मक्षेत्रे कुरुक्षेत्रे / समवेता युयुत्सवः / "
        "मामकाः पाण्डवाश्चैव / किमकुर्वत सञ्जय",
        "https://vedabase.io/en/library/bg/1/1/",
        "SAMPLE ROW — delete it and type your own anvyaya here.",
        "SAMPLE ROW — delete it and type your own translation here.",
    ],
}


def make(folder: str, name: str) -> None:
    path = os.path.join(ROOT, folder, f"{name}.xlsx")
    if os.path.exists(path) and not FORCE:
        print("kept   ", path)
        return
    os.makedirs(os.path.dirname(path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Verses"

    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="B5651D")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in TEXT_COLS:
        # Setting the column's own number_format makes Excel treat any new
        # text typed into that column as Text by default (no 2999-row loop
        # needed, which would only bloat the file).
        ws.column_dimensions[get_column_letter(col)].number_format = "@"

    sample = SAMPLES.get(name)
    if sample:
        for col, val in enumerate(sample, start=1):
            c = ws.cell(row=2, column=col, value=val)
            c.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    wb.save(path)
    print("wrote  ", path)


for name, _ in SB_FILES:
    make("sb", name)
for name, _ in BG_FILES:
    make("bg", name)
