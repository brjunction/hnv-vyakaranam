"""
One-off script used to generate the starter .xlsx files under content/chapters/.
You don't need to run this again — it's just here so you can see how the
template files were made, or regenerate a fresh blank one if you ever need to.

Usage: python3 scripts/make_templates.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "S.No",
    "Subchapter",
    "Sutra Number",
    "Sanskrit",
    "Transliteration",
    "Vrtti",
    "Translation",
]

CHAPTERS = {
    "mangalacarana": [],
    "samjna-sandhi-prakarana": [
        [1, "Sarvesvara-sandhi", "42", "अइउण्", "a i u Ṇ", "एतत् उदाहरणम्।", "Example vrtti text.", ],
    ],
    "nama-prakarana": [],
    "akhyata-prakarana": [],
    "karaka-prakarana": [],
    "krdanta-prakarana": [],
    "samasa-prakarana": [],
    "taddhita-prakarana": [],
    "afterword": [],
}

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "content", "chapters")

for slug, rows in CHAPTERS.items():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sutras"

    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="B5651D")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r, row in enumerate(rows, start=2):
        for col, val in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=val)

    widths = [6, 22, 14, 34, 28, 40, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    out_path = os.path.join(OUT_DIR, f"{slug}.xlsx")
    wb.save(out_path)
    print("wrote", out_path)
